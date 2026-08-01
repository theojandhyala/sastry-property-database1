#!/usr/bin/env python3
"""Build the SQLite property database from the landlord workbook.

    python3 scripts/import_xlsx.py \
        --xlsx data/source/landlord_10_property_database.xlsx \
        --db   data/property.db

The load is idempotent: the schema is dropped and recreated on every run, so
re-running against an updated workbook simply rebuilds the database.

Values the workbook computed with formulas are not imported. Where a formula
failed to evaluate (the Tenancies sheet stores XLOOKUP as the unsupported
'_xludf.XLOOKUP', which reads back as '#NAME?'), the importer resolves the
value from the referenced table instead and reports what it did.
"""

from __future__ import annotations

import argparse
import datetime as dt
import sqlite3
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency hint
    sys.exit("openpyxl is required: pip install -r requirements.txt")

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = REPO_ROOT / "data" / "source" / "landlord_10_property_database.xlsx"
DEFAULT_DB = REPO_ROOT / "data" / "property.db"

# Strings Excel leaves behind when a formula cannot be evaluated.
ERROR_VALUES = {"#NAME?", "#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NULL!", "#NUM!"}

notes: list[str] = []


def note(message: str) -> None:
    notes.append(message)


# ---------------------------------------------------------------------------
# Cell coercion
# ---------------------------------------------------------------------------
def clean(value):
    """Normalise a cell: blanks and unevaluated formula errors become None."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value in ERROR_VALUES:
            return None
    return value


def as_text(value):
    value = clean(value)
    return None if value is None else str(value)


def as_int(value):
    value = clean(value)
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def as_num(value):
    value = clean(value)
    if value is None:
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def as_date(value):
    """Return an ISO date string, accepting datetimes or text dates."""
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(str(value), fmt).date().isoformat()
        except ValueError:
            continue
    note(f"Unparsable date left empty: {value!r}")
    return None


def as_bool(value):
    """'Yes'/'No' (and friends) to 1/0."""
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    text = str(value).strip().lower()
    if text in {"yes", "y", "true", "1"}:
        return 1
    if text in {"no", "n", "false", "0"}:
        return 0
    return None


def rows(sheet):
    """Yield populated data rows (header skipped, trailing blanks dropped)."""
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and str(row[0]).strip():
            yield row


# ---------------------------------------------------------------------------
# Loaders, one per sheet
# ---------------------------------------------------------------------------
def load_lists(conn, wb):
    for col in wb["Lists"].iter_cols(values_only=True):
        category = as_text(col[0])
        if not category:
            continue
        order = 0
        for cell in col[1:]:
            value = as_text(cell)
            if value:
                order += 1
                conn.execute(
                    "INSERT OR IGNORE INTO enum_value (category, value, sort_order)"
                    " VALUES (?, ?, ?)",
                    (category, value, order),
                )


def load_landlords(conn, wb):
    for r in rows(wb["Landlords"]):
        conn.execute(
            "INSERT INTO landlord (landlord_id, name, phone, email, address, notes)"
            " VALUES (?, ?, ?, ?, ?, ?)",
            (as_text(r[0]), as_text(r[1]), as_text(r[2]), as_text(r[3]),
             as_text(r[4]), as_text(r[5])),
        )


def load_tenants(conn, wb):
    for r in rows(wb["Tenants"]):
        name = as_text(r[1])
        # The workbook parks unlet units against dummy tenants called 'Vacant'.
        placeholder = 1 if (name or "").lower() == "vacant" else 0
        conn.execute(
            "INSERT INTO tenant (tenant_id, name, phone, email, emergency_contact,"
            " notes, is_placeholder) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (as_text(r[0]), name, as_text(r[2]), as_text(r[3]), as_text(r[4]),
             as_text(r[5]), placeholder),
        )


def load_contractors(conn, wb):
    for r in rows(wb["Contractors"]):
        conn.execute(
            "INSERT INTO contractor (contractor_id, name, trade, phone, email, area,"
            " insurance_checked, notes) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (as_text(r[0]), as_text(r[1]), as_text(r[2]), as_text(r[3]),
             as_text(r[4]), as_text(r[5]), as_bool(r[6]), as_text(r[7])),
        )


def load_properties(conn, wb):
    """Load properties; return {property_id: mortgage_id} for cross-checking.

    'Annual Rent' is dropped - it is monthly_rent * 12 and lives in a view.
    """
    declared_mortgages = {}
    for r in rows(wb["Properties"]):
        property_id = as_text(r[0])
        conn.execute(
            "INSERT INTO property (property_id, landlord_id, address, town_city,"
            " postcode, property_type, bedrooms, status, monthly_rent, notes)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (property_id, as_text(r[1]), as_text(r[2]), as_text(r[3]), as_text(r[4]),
             as_text(r[5]), as_int(r[6]), as_text(r[7]), as_num(r[8]), as_text(r[11])),
        )
        if as_text(r[9]):
            declared_mortgages[property_id] = as_text(r[9])
    return declared_mortgages


def load_mortgages(conn, wb, declared_mortgages):
    for r in rows(wb["Mortgages"]):
        mortgage_id, property_id = as_text(r[0]), as_text(r[1])
        expected = declared_mortgages.get(property_id)
        if expected and expected != mortgage_id:
            note(f"Mortgage link mismatch for {property_id}: Properties sheet says "
                 f"{expected}, Mortgages sheet says {mortgage_id}")
        conn.execute(
            "INSERT INTO mortgage (mortgage_id, property_id, lender, balance,"
            " monthly_payment, interest_rate, fixed_until, notes)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (mortgage_id, property_id, as_text(r[2]), as_num(r[3]), as_num(r[4]),
             as_num(r[5]), as_date(r[6]), as_text(r[7])),
        )


def load_tenancies(conn, wb):
    """Load tenancies, repairing the broken XLOOKUP rent/deposit columns.

    Sheet formulas: rent = XLOOKUP(PropertyID -> Properties.Monthly Rent),
    deposit = the same cell. Both read back as '#NAME?', so recompute them.
    """
    rents = dict(conn.execute("SELECT property_id, monthly_rent FROM property"))
    repaired = 0
    for r in rows(wb["Tenancies"]):
        tenancy_id, property_id = as_text(r[0]), as_text(r[1])
        rent, deposit = as_num(r[5]), as_num(r[6])
        if rent is None:
            rent = rents.get(property_id)
            deposit = deposit if deposit is not None else rent
            repaired += 1
        conn.execute(
            "INSERT INTO tenancy (tenancy_id, property_id, tenant_id, start_date,"
            " end_date, monthly_rent, deposit, deposit_protected, status, notes)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (tenancy_id, property_id, as_text(r[2]), as_date(r[3]), as_date(r[4]),
             rent, deposit, as_bool(r[7]), as_text(r[8]), as_text(r[10])),
        )
    if repaired:
        note(f"Tenancies: rent/deposit resolved from the Properties sheet for "
             f"{repaired} row(s) whose XLOOKUP formulas had not been evaluated")


def load_rent_payments(conn, wb):
    """Arrears and the 'Mon yyyy' label are formulas; both are views here."""
    for r in rows(wb["Rent Payments"]):
        conn.execute(
            "INSERT INTO rent_payment (payment_id, property_id, tenant_id, due_date,"
            " paid_date, amount_due, amount_paid, status, notes)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (as_text(r[0]), as_text(r[1]), as_text(r[2]), as_date(r[3]), as_date(r[4]),
             as_num(r[5]) or 0, as_num(r[6]) or 0, as_text(r[7]), as_text(r[10])),
        )


def load_maintenance(conn, wb):
    for r in rows(wb["Maintenance"]):
        conn.execute(
            "INSERT INTO maintenance (maintenance_id, property_id, reported_date,"
            " description, contractor_id, priority, status, cost_estimate,"
            " actual_cost, completed_date, notes)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (as_text(r[0]), as_text(r[1]), as_date(r[2]), as_text(r[3]), as_text(r[4]),
             as_text(r[5]), as_text(r[6]), as_num(r[7]), as_num(r[8]), as_date(r[9]),
             as_text(r[10])),
        )


def load_certificates(conn, wb):
    for r in rows(wb["Certificates"]):
        conn.execute(
            "INSERT INTO certificate (certificate_id, property_id, certificate_type,"
            " issue_date, expiry_date, status, provider, cost, document_link, notes)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (as_text(r[0]), as_text(r[1]), as_text(r[2]), as_date(r[3]), as_date(r[4]),
             as_text(r[5]), as_text(r[6]), as_num(r[7]), as_text(r[8]), as_text(r[9])),
        )


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------
TABLES = ["enum_value", "landlord", "tenant", "contractor", "property", "mortgage",
          "tenancy", "rent_payment", "maintenance", "certificate"]


def build(xlsx_path: Path, db_path: Path) -> None:
    if not xlsx_path.exists():
        sys.exit(f"Workbook not found: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=False)
    db_path.parent.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        for script in ("schema.sql", "views.sql"):
            conn.executescript((REPO_ROOT / "sql" / script).read_text())

        load_lists(conn, wb)
        load_landlords(conn, wb)
        load_tenants(conn, wb)
        load_contractors(conn, wb)
        declared_mortgages = load_properties(conn, wb)
        load_mortgages(conn, wb, declared_mortgages)
        load_tenancies(conn, wb)
        load_rent_payments(conn, wb)
        load_maintenance(conn, wb)
        load_certificates(conn, wb)

        violations = conn.execute("PRAGMA foreign_key_check").fetchall()
        if violations:
            raise SystemExit(f"Foreign key violations, load rolled back: {violations}")
        conn.commit()

        print(f"Loaded {xlsx_path.name} -> {db_path}\n")
        print("Rows")
        for table in TABLES:
            count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            print(f"  {table:<14} {count:>5}")

        print("\nDashboard")
        for metric, value in conn.execute(
            "SELECT metric, value FROM v_portfolio_dashboard ORDER BY sort_order"
        ):
            print(f"  {metric:<28} {value:>12,.2f}")

        issues = conn.execute(
            "SELECT entity, record_id, issue FROM v_data_quality_issues"
            " ORDER BY entity, record_id"
        ).fetchall()
        if notes or issues:
            print("\nData quality")
            for message in notes:
                print(f"  [import] {message}")
            for entity, record_id, issue in issues:
                print(f"  [{entity}] {record_id}: {issue}")
    finally:
        conn.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX,
                        help="source workbook (default: %(default)s)")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB,
                        help="SQLite file to build (default: %(default)s)")
    args = parser.parse_args()
    build(args.xlsx, args.db)


if __name__ == "__main__":
    main()
